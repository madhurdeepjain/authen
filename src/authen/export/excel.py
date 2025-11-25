"""
Excel export for validated references.

Exports references to Excel format with:
- All reference fields as columns
- Authors serialized as JSON
- Validation status and confidence
- Optional raw OpenAlex data
"""

import json
from pathlib import Path

import pandas as pd
import structlog

from authen.core.schemas import (
    PipelineResult,
    ReferenceData,
    ValidationResult,
)

logger = structlog.get_logger()


class ExcelExporter:
    """
    Export validated references to Excel format.

    Authors are serialized as JSON in a single column to preserve
    all author details including affiliations.
    """

    def __init__(
        self,
        include_raw_openalex: bool = False,
        include_validation_details: bool = True,
    ):
        """
        Initialize the exporter.

        Args:
            include_raw_openalex: Include raw OpenAlex response
            include_validation_details: Include validation status and scores
        """
        self.include_raw_openalex = include_raw_openalex
        self.include_validation_details = include_validation_details

    def export(
        self,
        data: list[ValidationResult] | PipelineResult,
        output_path: str | Path,
    ) -> Path:
        """
        Export references to Excel.

        Args:
            data: Validation results or pipeline result
            output_path: Path for the output Excel file

        Returns:
            Path to the created file
        """
        output_path = Path(output_path)

        # Extract validation results
        if isinstance(data, PipelineResult):
            results = data.validation_results
        else:
            results = data

        logger.info("exporting_to_excel", count=len(results), path=str(output_path))

        # Build rows
        rows = []
        for i, result in enumerate(results):
            row = self._build_row(result, i + 1)
            rows.append(row)

        # Create DataFrame
        df = pd.DataFrame(rows)

        # Reorder columns
        df = self._reorder_columns(df)

        # Write to Excel with formatting
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="References", index=False)

            # Get worksheet for formatting
            worksheet = writer.sheets["References"]

            # Apply formatting
            self._apply_formatting(worksheet, df)

        logger.info("export_complete", path=str(output_path))
        return output_path

    def _build_row(self, result: ValidationResult, index: int) -> dict:
        """Build a row dictionary from a validation result."""
        # Use validated data if available, otherwise original
        ref = result.get_best_reference()

        row = {
            "reference_number": ref.reference_number or index,
            "title": ref.title,
            "authors_json": self._serialize_authors(ref),
            "year": ref.year,
            "publication": ref.publication,
            "publisher": ref.publisher,
            "volume": ref.volume,
            "issue": ref.issue,
            "pages": ref.pages,
            "doi": ref.doi,
            "isbn": ref.isbn,
            "pmid": ref.pmid,
            "arxiv_id": ref.arxiv_id,
            "url": ref.url,
            "work_type": ref.work_type,
            "openalex_id": ref.openalex_id,
            "cited_by_count": ref.cited_by_count,
            "is_open_access": ref.is_open_access,
            "raw_text": ref.raw_text,
        }

        # Add validation details
        if self.include_validation_details:
            row.update(
                {
                    "validation_status": result.status.value,
                    "validation_confidence": round(result.confidence, 3),
                    "title_similarity": (
                        round(result.title_similarity, 3)
                        if result.title_similarity
                        else None
                    ),
                    "author_similarity": (
                        round(result.author_similarity, 3)
                        if result.author_similarity
                        else None
                    ),
                    "match_method": result.match_method,
                    "openalex_url": result.openalex_url,
                    "error_message": result.error_message,
                }
            )

        # Add raw OpenAlex data
        if self.include_raw_openalex and result.openalex_raw:
            row["openalex_raw_json"] = json.dumps(result.openalex_raw)

        return row

    def _serialize_authors(self, ref: ReferenceData) -> str:
        """Serialize authors list to JSON."""
        if not ref.authors:
            return "[]"

        authors_data = []
        for author in ref.authors:
            author_dict = {
                "first_name": author.first_name,
                "last_name": author.last_name,
                "full_name": author.full_name or author.display_name,
            }

            # Add optional fields if present
            if author.orcid:
                author_dict["orcid"] = author.orcid
            if author.openalex_id:
                author_dict["openalex_id"] = author.openalex_id
            if author.emails:
                author_dict["emails"] = author.emails

            # Add affiliations
            if author.affiliations:
                author_dict["affiliations"] = [
                    {
                        "name": aff.name,
                        "department": aff.department,
                        "country": aff.country,
                        "city": aff.city,
                        "ror_id": aff.ror_id,
                        "openalex_id": aff.openalex_id,
                    }
                    for aff in author.affiliations
                    if aff.name  # Only include affiliations with names
                ]

            # Clean None values
            author_dict = {k: v for k, v in author_dict.items() if v is not None}

            authors_data.append(author_dict)

        return json.dumps(authors_data, ensure_ascii=False)

    def _reorder_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        """Reorder columns for better readability."""
        preferred_order = [
            "reference_number",
            "title",
            "authors_json",
            "year",
            "publication",
            "publisher",
            "volume",
            "issue",
            "pages",
            "doi",
            "url",
            "work_type",
            "validation_status",
            "validation_confidence",
            "title_similarity",
            "author_similarity",
            "match_method",
            "openalex_id",
            "openalex_url",
            "cited_by_count",
            "is_open_access",
            "isbn",
            "pmid",
            "arxiv_id",
            "error_message",
            "raw_text",
            "openalex_raw_json",
        ]

        # Filter to columns that exist
        existing = [c for c in preferred_order if c in df.columns]
        # Add any remaining columns
        remaining = [c for c in df.columns if c not in existing]

        return df[existing + remaining]

    def _apply_formatting(self, worksheet, df: pd.DataFrame) -> None:
        """Apply formatting to the worksheet."""
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        # Header styling
        header_font = Font(bold=True)
        header_fill = PatternFill(
            start_color="CCE5FF", end_color="CCE5FF", fill_type="solid"
        )

        for col_num, column_title in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill

        # Set column widths
        column_widths = {
            "reference_number": 12,
            "title": 50,
            "authors_json": 40,
            "year": 8,
            "publication": 30,
            "publisher": 20,
            "doi": 30,
            "validation_status": 15,
            "validation_confidence": 15,
            "raw_text": 50,
        }

        for col_num, column_title in enumerate(df.columns, 1):
            col_letter = get_column_letter(col_num)
            width = column_widths.get(column_title, 15)
            worksheet.column_dimensions[col_letter].width = width

        # Wrap text for long columns
        wrap_columns = ["title", "authors_json", "raw_text", "openalex_raw_json"]
        for col_num, column_title in enumerate(df.columns, 1):
            if column_title in wrap_columns:
                for row in range(2, len(df) + 2):
                    cell = worksheet.cell(row=row, column=col_num)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Color code validation status
        status_colors = {
            "validated": "C6EFCE",  # Light green
            "partial_match": "FFEB9C",  # Light yellow
            "not_found": "FFC7CE",  # Light red
            "error": "FFC7CE",  # Light red
            "pending": "E0E0E0",  # Light gray
        }

        if "validation_status" in df.columns:
            status_col = df.columns.get_loc("validation_status") + 1
            for row in range(2, len(df) + 2):
                cell = worksheet.cell(row=row, column=status_col)
                status = cell.value
                if status in status_colors:
                    cell.fill = PatternFill(
                        start_color=status_colors[status],
                        end_color=status_colors[status],
                        fill_type="solid",
                    )


def export_to_excel(
    results: list[ValidationResult] | PipelineResult,
    output_path: str | Path,
    include_raw_openalex: bool = False,
) -> Path:
    """
    Convenience function to export to Excel.

    Args:
        results: Validation results or pipeline result
        output_path: Output file path
        include_raw_openalex: Include raw API response

    Returns:
        Path to created file
    """
    exporter = ExcelExporter(include_raw_openalex=include_raw_openalex)
    return exporter.export(results, output_path)


def export_to_json(
    results: list[ValidationResult] | PipelineResult,
    output_path: str | Path,
) -> Path:
    """
    Export results to JSON.

    Args:
        results: Validation results or pipeline result
        output_path: Output file path

    Returns:
        Path to created file
    """
    output_path = Path(output_path)

    if isinstance(results, PipelineResult):
        data = results.model_dump()
    else:
        data = [r.model_dump() for r in results]

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=str)

    logger.info("exported_to_json", path=str(output_path))
    return output_path


def export_to_csv(
    results: list[ValidationResult] | PipelineResult,
    output_path: str | Path,
) -> Path:
    """
    Export results to CSV (simplified format).

    Args:
        results: Validation results or pipeline result
        output_path: Output file path

    Returns:
        Path to created file
    """
    output_path = Path(output_path)

    if isinstance(results, PipelineResult):
        validation_results = results.validation_results
    else:
        validation_results = results

    rows = []
    for i, result in enumerate(validation_results):
        ref = result.get_best_reference()
        row = {
            "reference_number": ref.reference_number or i + 1,
            "title": ref.title,
            "authors": ", ".join(a.display_name for a in ref.authors),
            "year": ref.year,
            "publication": ref.publication,
            "doi": ref.doi,
            "validation_status": result.status.value,
            "confidence": result.confidence,
        }
        rows.append(row)

    df = pd.DataFrame(rows)
    df.to_csv(output_path, index=False)

    logger.info("exported_to_csv", path=str(output_path))
    return output_path
